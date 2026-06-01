"""Stage 6 — Write Excel: the five-column register, ready to hand to a PM.

Plain openpyxl. The columns match the build spec exactly. A little formatting
goes a long way in a demo: bold frozen header, sensible widths, and Urgent rows
tinted so the eye lands on the schedule risks first.
"""
from dataclasses import dataclass

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = ["Spec Section", "Item", "Submittal Required", "Lead Time Category", "Priority Flag"]
_WIDTHS = [34, 40, 40, 22, 14]

_HEADER_FILL = PatternFill("solid", fgColor="1F2937")  # slate
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_URGENT_FILL = PatternFill("solid", fgColor="FDE2E1")  # soft red
_URGENT_FONT = Font(color="991B1B", bold=True)


@dataclass
class Row:
    spec_section: str
    item: str
    submittal_required: str
    lead_time_category: str
    priority: str


def write_register(rows: list[Row], out_path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Submittal Register"

    ws.append(HEADERS)
    for col, width in enumerate(_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
        cell = ws.cell(row=1, column=col)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center")

    for r in rows:
        ws.append([r.spec_section, r.item, r.submittal_required, r.lead_time_category, r.priority])
        excel_row = ws.max_row
        for col in range(1, len(HEADERS) + 1):
            ws.cell(row=excel_row, column=col).alignment = Alignment(
                vertical="top", wrap_text=True
            )
        if r.priority == "Urgent":
            for col in range(1, len(HEADERS) + 1):
                ws.cell(row=excel_row, column=col).fill = _URGENT_FILL
            ws.cell(row=excel_row, column=5).font = _URGENT_FONT

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{ws.max_row}"
    wb.save(out_path)
