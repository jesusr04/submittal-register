"""Offline tests for the deterministic logic: lead-time lookup, priority, Excel.

No API key needed. These lock in the rules a reviewer will scrutinize on a demo.
"""
import os
import tempfile
import unittest

from openpyxl import load_workbook

from src.excel import HEADERS, Row, write_register
from src.lead_times import CATEGORIES, bucket_for, cell_text
from src.priority import priority_flag


class TestLeadTimes(unittest.TestCase):
    def test_every_category_resolves(self):
        for cat in CATEGORIES:
            label, low, high = bucket_for(cat)
            self.assertIn(label, {"Stock", "Short", "Medium", "Long", "Extra-Long"})
            self.assertLessEqual(low, high)

    def test_unknown_is_conservative_medium(self):
        self.assertEqual(bucket_for("not-a-real-category")[0], "Medium")

    def test_cell_text_formats(self):
        self.assertEqual(cell_text("concrete_rebar_aggregate"), "Stock")
        self.assertEqual(cell_text("storefront_glazing"), "Long (20-28 wk)")


class TestPriority(unittest.TestCase):
    def test_stock_is_low(self):
        self.assertEqual(priority_flag("Stock"), "Low")

    def test_plan_is_medium(self):
        self.assertEqual(priority_flag("Plan"), "Medium")

    def test_short_lead_is_medium(self):
        self.assertEqual(priority_flag("Short"), "Medium")

    def test_medium_lead_is_high(self):
        self.assertEqual(priority_flag("Medium"), "High")

    def test_long_and_extra_long_are_urgent(self):
        self.assertEqual(priority_flag("Long"), "Urgent")
        self.assertEqual(priority_flag("Extra-Long"), "Urgent")

    def test_unknown_bucket_surfaces_as_high(self):
        self.assertEqual(priority_flag("not-a-bucket"), "High")


class TestExcel(unittest.TestCase):
    def test_writes_readable_register(self):
        rows = [
            Row("26 24 13 — Switchgear", "Switchgear",
                "Shop drawings; product data", "Extra-Long (30-52 wk)", "Urgent"),
            Row("03 30 00 — Cast-in-Place Concrete", "Concrete design mix",
                "Design mixtures; product data", "Stock", "Low"),
            Row("03 30 00 — Cast-in-Place Concrete", "Curing method",
                "Proposed curing method", "Plan", "Medium"),
        ]
        with tempfile.TemporaryDirectory() as d:
            out = os.path.join(d, "register.xlsx")
            write_register(rows, out)
            self.assertTrue(os.path.exists(out))
            ws = load_workbook(out).active
            self.assertEqual([c.value for c in ws[1]], HEADERS)
            self.assertEqual(ws.max_row, 4)  # header + 3 rows
            self.assertEqual(ws.cell(row=2, column=5).value, "Urgent")
            self.assertEqual(ws.cell(row=4, column=4).value, "Plan")


if __name__ == "__main__":
    unittest.main()
